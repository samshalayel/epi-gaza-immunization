from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Monthly Tracker"

MONTHS  = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
GOVS    = ["North Gaza","Gaza City","Deir Al-Balah","Khan Younis","Rafah"]
TARGETS = [85000, 180000, 120000, 160000, 75000]

def col(n):
    return get_column_letter(n)

# ── Row 1: Title ──────────────────────────────────────────────────
ws.merge_cells("A1:" + col(28) + "1")
c = ws["A1"]
c.value     = "Gaza Strip EPI  –  Monthly Vaccination Tracker 2026"
c.font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
c.fill      = PatternFill("solid", fgColor="1F4E79")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 28

# ── Rows 2-3: Column headers ──────────────────────────────────────
def hdr_cell(cell_ref, value, bg="1F4E79", size=10, wrap=False):
    c = ws[cell_ref]
    c.value     = value
    c.font      = Font(name="Arial", bold=True, size=size, color="FFFFFF")
    c.fill      = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    return c

ws.merge_cells("A2:A3"); hdr_cell("A2", "Governorate")
ws.merge_cells("B2:B3"); hdr_cell("B2", "Annual\nTarget", wrap=True)

BLUES = ["1F4E79", "2E75B6"]
for i, month in enumerate(MONTHS):
    mc = 3 + i * 2
    cc = 4 + i * 2
    h1 = BLUES[i % 2]
    h2 = BLUES[1 - i % 2]
    ws.merge_cells(col(mc) + "2:" + col(cc) + "2")
    hdr_cell(col(mc) + "2", month, bg=h1)
    hdr_cell(col(mc) + "3", "Monthly",    bg=h1, size=9)
    hdr_cell(col(cc) + "3", "Cumulative", bg=h2, size=9)

ws.merge_cells(col(27) + "2:" + col(27) + "3"); hdr_cell(col(27) + "2", "Annual\nTotal", wrap=True)
ws.merge_cells(col(28) + "2:" + col(28) + "3"); hdr_cell(col(28) + "2", "Coverage\n%",   wrap=True)
ws.row_dimensions[2].height = 22
ws.row_dimensions[3].height = 16

# ── Data rows (4-8) ───────────────────────────────────────────────
DATA_START = 4
for r, (gov, target) in enumerate(zip(GOVS, TARGETS)):
    row    = DATA_START + r
    row_bg = "FFFFFF" if r % 2 == 0 else "EEF5FB"

    c = ws.cell(row=row, column=1, value=gov)
    c.font = Font(name="Arial", bold=True, size=10)
    c.fill = PatternFill("solid", fgColor=row_bg)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    c = ws.cell(row=row, column=2, value=target)
    c.font = Font(name="Arial", size=10, color="0000FF")
    c.fill = PatternFill("solid", fgColor=row_bg)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.number_format = "#,##0"

    for i in range(12):
        mc = 3 + i * 2
        cc = 4 + i * 2

        # Leave blank by default — blank = not yet entered, 0 = real zero
        inp = ws.cell(row=row, column=mc, value=None)
        inp.font          = Font(name="Arial", size=10, color="0000FF")
        inp.fill          = PatternFill("solid", fgColor=row_bg)
        inp.alignment     = Alignment(horizontal="right")
        inp.number_format = "#,##0"

        # Cumulative stops at last entered month:
        # - If this month is blank → blank (stop)
        # - If this month has data and prev cumul exists → prev + this
        # - If this month has data but prev cumul is blank → just this month
        cum = ws.cell(row=row, column=cc)
        mc_ref = col(mc) + str(row)
        if i == 0:
            cum.value = '=IF(ISBLANK(' + mc_ref + '),"",(' + mc_ref + '))'
        else:
            prev_cc_ref = col(4 + (i - 1) * 2) + str(row)
            cum.value = (
                '=IF(ISBLANK(' + mc_ref + '),"",IF(' + prev_cc_ref + '="",'
                + mc_ref + ',' + prev_cc_ref + '+' + mc_ref + '))'
            )
        cum.font          = Font(name="Arial", size=10, bold=True, color="000000")
        cum.fill          = PatternFill("solid", fgColor="D6E4F0")
        cum.alignment     = Alignment(horizontal="right")
        cum.number_format = "#,##0"

    # Annual Total = SUM of monthly cols only (SUM ignores blanks — correct)
    monthly_cols = ",".join(col(3 + i * 2) + str(row) for i in range(12))
    a = ws.cell(row=row, column=27)
    a.value         = "=SUM(" + monthly_cols + ")"
    a.font          = Font(name="Arial", size=10, bold=True)
    a.fill          = PatternFill("solid", fgColor="BDD7EE")
    a.alignment     = Alignment(horizontal="right")
    a.number_format = "#,##0"

    v = ws.cell(row=row, column=28)
    v.value         = "=IF(B" + str(row) + ">0," + col(27) + str(row) + "/B" + str(row) + ",\"\")"
    v.font          = Font(name="Arial", size=10, bold=True)
    v.fill          = PatternFill("solid", fgColor="BDD7EE")
    v.alignment     = Alignment(horizontal="center")
    v.number_format = "0.0%"
    ws.row_dimensions[row].height = 19

# ── Total row ─────────────────────────────────────────────────────
total_row = DATA_START + len(GOVS)

def tot(col_n, formula, fmt="#,##0", cumul=False):
    c = ws.cell(row=total_row, column=col_n)
    c.value         = formula
    c.font          = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    c.fill          = PatternFill("solid", fgColor="2E75B6" if cumul else "1F4E79")
    c.alignment     = Alignment(horizontal="center" if fmt == "0.0%" else "right", vertical="center")
    c.number_format = fmt

c = ws.cell(row=total_row, column=1, value="TOTAL")
c.font      = Font(name="Arial", bold=True, size=10, color="FFFFFF")
c.fill      = PatternFill("solid", fgColor="1F4E79")
c.alignment = Alignment(horizontal="left", indent=1, vertical="center")
ws.row_dimensions[total_row].height = 21

tot(2, "=SUM(B" + str(DATA_START) + ":B" + str(total_row - 1) + ")")
for i in range(12):
    mc = 3 + i * 2
    cc = 4 + i * 2
    tot(mc, "=SUM(" + col(mc) + str(DATA_START) + ":" + col(mc) + str(total_row - 1) + ")")
    # Cumulative total: sum only rows where cumul is not blank
    tot(cc,
        "=SUMIF(" + col(cc) + str(DATA_START) + ":" + col(cc) + str(total_row - 1) + ',\"<>\"\"\")',
        cumul=True)
tot(27, "=SUM(" + col(27) + str(DATA_START) + ":" + col(27) + str(total_row - 1) + ")")
tot(28, "=IF(B" + str(total_row) + ">0," + col(27) + str(total_row) + "/B" + str(total_row) + ",\"\")", fmt="0.0%")

# ── Legend ────────────────────────────────────────────────────────
leg = total_row + 2
ws.merge_cells("A" + str(leg) + ":" + col(28) + str(leg))
c = ws.cell(row=leg, column=1)
c.value     = "Blue = enter monthly doses (leave BLANK if month not yet reported — do NOT enter 0)   |   Light blue = cumulative stops at last entered month automatically"
c.font      = Font(name="Arial", size=9, italic=True, color="555555")
c.alignment = Alignment(horizontal="left", vertical="center")

# ── Column widths & freeze ────────────────────────────────────────
ws.column_dimensions["A"].width = 15
ws.column_dimensions["B"].width = 10
for i in range(12):
    ws.column_dimensions[col(3 + i * 2)].width = 9
    ws.column_dimensions[col(4 + i * 2)].width = 10
ws.column_dimensions[col(27)].width = 10
ws.column_dimensions[col(28)].width = 9

ws.freeze_panes = "C4"

OUT = r"C:\Users\Administrator\Desktop\Gaza_EPI_Monthly_Tracker.xlsx"
wb.save(OUT)
print("Saved:", OUT)
